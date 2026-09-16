"""IEEE conference formatting for python-docx.

Follows the IEEE conference Word template: US Letter, Times New Roman, a
single-column band for the title and abstract, then a two-column body.
Measurements are from the template's own specification rather than eyeballed
from a PDF.

python-docx has no API for multi-column sections, so :func:`set_columns`
edits the section's ``<w:cols>`` element directly. That is the documented
workaround, not a hack around a bug.
"""

from __future__ import annotations

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"
MONO = "Consolas"

TITLE_PT = 20      # template says 24; 20 keeps this long title off a 4th line
AUTHOR_PT = 11
ABSTRACT_PT = 9
BODY_PT = 10
CAPTION_PT = 8
REF_PT = 8

COL_GAP_IN = 0.25


def set_columns(section, n: int, gap_in: float = COL_GAP_IN) -> None:
    """Set the column count on a section."""
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(gap_in * 1440)))
    cols.set(qn("w:equalWidth"), "1")


def page_setup(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)


def new_section(doc, cols: int):
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    page_setup(s)
    set_columns(s, cols)
    return s


def _style_run(run, *, size, bold=False, italic=False, mono=False,
               caps=False) -> None:
    run.font.name = MONO if mono else FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if caps:
        run.font.small_caps = True
    # East-Asian font must be set too or Word substitutes for some glyphs.
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), MONO if mono else FONT)


def para(doc, *, align=None, space_after=0, space_before=0, indent=0.0,
         keep_with_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.0
    if indent:
        pf.first_line_indent = Inches(indent)
    if align is not None:
        p.alignment = align
    if keep_with_next:
        pf.keep_with_next = True
    return p


def emit_runs(p, chunks, *, size=BODY_PT, base_bold=False, base_italic=False):
    """Write ``(text, style)`` pairs from tex_blocks.runs into a paragraph."""
    for text, style in chunks:
        r = p.add_run(text)
        _style_run(
            r,
            size=size,
            bold=base_bold or style == "b",
            italic=base_italic or style == "i",
            mono=style == "tt",
        )


def heading(doc, numeral: str, text: str):
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12,
             space_after=4, keep_with_next=True)
    r = p.add_run(f"{numeral}.  {text}")
    _style_run(r, size=BODY_PT, caps=True)
    return p


def subheading(doc, letter: str, text: str):
    p = para(doc, space_before=8, space_after=3, keep_with_next=True)
    r = p.add_run(f"{letter}. {text}")
    _style_run(r, size=BODY_PT, italic=True)
    return p


def title_block(doc, title: str, author_lines: list[str]) -> None:
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    r = p.add_run(title)
    _style_run(r, size=TITLE_PT)
    for i, line in enumerate(author_lines):
        q = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_after=2 if i < len(author_lines) - 1 else 12)
        rr = q.add_run(line)
        _style_run(rr, size=AUTHOR_PT)


def mono_block(doc, lines: list[str], *, size=7.5) -> None:
    """A fixed-width block, used for the small tables IEEE sets as figures."""
    for line in lines:
        p = para(doc, space_after=0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        _style_run(r, size=size, mono=True)


def table(doc, rows: list[list[str]], *, size=7.5):
    if not rows:
        return None
    width = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=width)
    t.style = "Table Grid"
    t.autofit = True
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(width):
            txt = row[ci] if ci < len(row) else ""
            cp = cells[ci].paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            r = cp.add_run(txt)
            _style_run(r, size=size, bold=(ri == 0))
    return t


def caption(doc, text: str) -> None:
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4,
             space_after=8)
    r = p.add_run(text)
    _style_run(r, size=CAPTION_PT)


def caption_runs(doc, label: str, chunks) -> None:
    """An IEEE caption: bold label, then the caption text at caption size."""
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4,
             space_after=8)
    r = p.add_run(label)
    _style_run(r, size=CAPTION_PT, bold=True)
    emit_runs(p, chunks, size=CAPTION_PT)
